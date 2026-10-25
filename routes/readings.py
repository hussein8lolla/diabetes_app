import statistics
from datetime import datetime, timedelta
from io import BytesIO

from flask import (Blueprint, render_template, redirect, url_for,
                   flash, request, send_file)
from flask_login import login_required, current_user

from models import db, Reading
from forms import ReadingForm

readings_bp = Blueprint('readings', __name__, url_prefix='/readings')


@readings_bp.route('/add', methods=['GET', 'POST'])
@login_required
def add():
    form = ReadingForm()
    if request.method == 'GET':
        form.date.data = datetime.now().date()
        form.time.data = datetime.now().time().replace(second=0, microsecond=0)

    if form.validate_on_submit():
        reading = Reading(
            user_id=current_user.id,
            date=form.date.data.strftime('%Y-%m-%d'),
            time=form.time.data.strftime('%H:%M'),
            value=form.value.data,
            measurement_time=form.measurement_time.data,
            medication_name=form.medication_name.data or '',
            medication_dose=form.medication_dose.data or 0.0,
            carbs=form.carbs.data or 0,  # ✅ تمت الإضافة
            note=form.note.data or ''
        )
        db.session.add(reading)
        db.session.commit()
        flash('✅ تم حفظ القياس بنجاح!', 'success')
        return redirect(url_for('main.dashboard'))

    return render_template('add_reading.html', form=form)


@readings_bp.route('/edit/<int:reading_id>', methods=['GET', 'POST'])
@login_required
def edit(reading_id):
    reading = Reading.query.filter_by(id=reading_id, user_id=current_user.id).first_or_404()
    form = ReadingForm(obj=reading)

    if request.method == 'GET':
        form.date.data = datetime.strptime(reading.date, '%Y-%m-%d').date()
        form.time.data = datetime.strptime(reading.time, '%H:%M').time()

    if form.validate_on_submit():
        reading.date = form.date.data.strftime('%Y-%m-%d')
        reading.time = form.time.data.strftime('%H:%M')
        reading.value = form.value.data
        reading.measurement_time = form.measurement_time.data
        reading.medication_name = form.medication_name.data or ''
        reading.medication_dose = form.medication_dose.data or 0.0
        reading.carbs = form.carbs.data or 0  # ✅ تمت الإضافة
        reading.note = form.note.data or ''
        db.session.commit()
        flash('✅ تم تحديث القياس!', 'success')
        return redirect(url_for('main.dashboard'))

    return render_template('edit_reading.html', form=form, reading=reading)


@readings_bp.route('/delete/<int:reading_id>', methods=['POST'])
@login_required
def delete(reading_id):
    reading = Reading.query.filter_by(id=reading_id, user_id=current_user.id).first_or_404()
    db.session.delete(reading)
    db.session.commit()
    flash('🗑️ تم حذف القياس.', 'info')
    return redirect(url_for('main.dashboard'))


@readings_bp.route('/clear-all', methods=['POST'])
@login_required
def clear_all():
    Reading.query.filter_by(user_id=current_user.id).delete()
    db.session.commit()
    flash('🗑️ تم مسح جميع القياسات.', 'warning')
    return redirect(url_for('main.dashboard'))


@readings_bp.route('/stats')
@login_required
def stats():
    def get_vals(days):
        start = (datetime.now() - timedelta(days=days)).strftime('%Y-%m-%d')
        rows = Reading.query.filter_by(user_id=current_user.id) \
            .filter(Reading.date >= start).all()
        return [r.value for r in rows]

    today_vals = get_vals(0)
    week_vals = get_vals(7)
    month_vals = get_vals(30)

    avg_today = round(statistics.mean(today_vals), 1) if today_vals else 0
    avg_week = round(statistics.mean(week_vals), 1) if week_vals else 0
    avg_month = round(statistics.mean(month_vals), 1) if month_vals else 0
    std_week = round(statistics.stdev(week_vals), 1) if len(week_vals) > 1 else 0
    std_month = round(statistics.stdev(month_vals), 1) if len(month_vals) > 1 else 0

    # ================= حساب TIR و eA1c =================
    total_readings = len(month_vals)
    if total_readings > 0:
        in_range = sum(1 for v in month_vals if 80 <= v <= 180)
        high = sum(1 for v in month_vals if v > 180)
        low = sum(1 for v in month_vals if v < 80)

        tir_percentage = round((in_range / total_readings) * 100)
        high_percentage = round((high / total_readings) * 100)
        low_percentage = round((low / total_readings) * 100)

        # حساب eA1c: (المتوسط + 46.7) / 28.7
        ea1c = round((avg_month + 46.7) / 28.7, 1)

        # تحديد حالة eA1c
        if ea1c < 5.7:
            a1c_status = "success"
            a1c_text = "طبيعي"
        elif ea1c < 6.5:
            a1c_status = "warning"
            a1c_text = "متوسط (ما قبل السكري)"
        else:
            a1c_status = "danger"
            a1c_text = "مرتفع"
    else:
        tir_percentage = high_percentage = low_percentage = 0
        ea1c = 0
        a1c_status = "secondary"
        a1c_text = "لا توجد قراءات كافية"

    # إحصائيات وقت القياس
    time_labels = [
        ('Morning', 'الصبح'), ('Noon', 'نص اليوم'), ('Afternoon', 'بعد الظهر'),
        ('Evening', 'بالليل'), ('Night', 'بليل'),
        ('Before Breakfast', 'قبل الفطار'), ('After Breakfast (2h)', 'بعد الفطار بساعتين'),
        ('Before Lunch', 'قبل الغداء'), ('After Lunch (2h)', 'بعد الغداء بساعتين'),
        ('Before Dinner', 'قبل العشاء'), ('After Dinner (2h)', 'بعد العشاء بساعتين'),
        ('Before Exercise', 'قبل التمرين'), ('After Exercise', 'بعد التمرين'),
        ('Before Sleep', 'قبل النوم'),
    ]
    time_stats = []
    for en, ar in time_labels:
        start = (datetime.now() - timedelta(days=30)).strftime('%Y-%m-%d')
        rows = Reading.query.filter_by(user_id=current_user.id, measurement_time=en) \
            .filter(Reading.date >= start).all()
        values = [r.value for r in rows]
        avg = round(statistics.mean(values), 1) if values else 0
        time_stats.append({'name_ar': ar, 'avg': avg, 'count': len(values)})

    chart_start = (datetime.now() - timedelta(days=30)).strftime('%Y-%m-%d')
    chart_rows = Reading.query.filter_by(user_id=current_user.id) \
        .filter(Reading.date >= chart_start) \
        .order_by(Reading.date.asc(), Reading.time.asc()).all()
    chart_labels = [f"{r.date} {r.time}" for r in chart_rows]
    chart_values = [r.value for r in chart_rows]

    return render_template('stats.html',
                           avg_today=avg_today, avg_week=avg_week, avg_month=avg_month,
                           std_week=std_week, std_month=std_month, time_stats=time_stats,
                           chart_labels=chart_labels, chart_values=chart_values,
                           tir_percentage=tir_percentage, high_percentage=high_percentage,
                           low_percentage=low_percentage,
                           ea1c=ea1c, a1c_status=a1c_status, a1c_text=a1c_text, total_readings=total_readings
                           )


@readings_bp.route('/export/excel')
@login_required
def export_excel():
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Diabetes Readings"

    headers = ["Date", "Time", "Value", "Measurement Time", "Medication", "Dose", "Carbs (g)", "Notes"]  # ✅ تمت الإضافة
    ws.append(headers)

    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="1A73E8", end_color="1A73E8", fill_type="solid")
    for col in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")

    readings = Reading.query.filter_by(user_id=current_user.id) \
        .order_by(Reading.date.desc(), Reading.time.desc()).all()
    for r in readings:
        ws.append([r.date, r.time, r.value, r.measurement_time,
                   r.medication_name, r.medication_dose, r.carbs, r.note])  # ✅ تمت الإضافة

    for col in ws.columns:
        max_len = max((len(str(cell.value)) if cell.value else 0) for cell in col)
        ws.column_dimensions[col[0].column_letter].width = max_len + 5

    output = BytesIO()
    wb.save(output)
    output.seek(0)

    filename = f"diabetes_report_{current_user.username}_{datetime.now().strftime('%Y%m%d')}.xlsx"
    return send_file(output,
                     mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
                     as_attachment=True, download_name=filename)


@readings_bp.route('/export/pdf')
@login_required
def export_pdf():
    from reportlab.lib.pagesizes import letter
    from reportlab.lib import colors
    from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib.units import inch

    output = BytesIO()
    doc = SimpleDocTemplate(output, pagesize=letter,
                            rightMargin=30, leftMargin=30,
                            topMargin=30, bottomMargin=30)

    styles = getSampleStyleSheet()
    title_style = ParagraphStyle('TitleStyle', parent=styles['Title'],
                                 fontSize=18, textColor=colors.HexColor('#1a73e8'),
                                 alignment=1, spaceAfter=20)
    subtitle_style = ParagraphStyle('SubtitleStyle', parent=styles['Heading2'],
                                    fontSize=13, textColor=colors.HexColor('#333'), spaceAfter=10)

    elements = []
    elements.append(Paragraph("Diabetes Readings Report", title_style))
    elements.append(Paragraph(f"User: {current_user.username}", styles['Normal']))
    elements.append(Paragraph(f"Email: {current_user.email}", styles['Normal']))
    elements.append(Paragraph(f"Generated: {datetime.now().strftime('%Y-%m-%d %H:%M')}", styles['Normal']))
    elements.append(Spacer(1, 20))

    def get_vals(days):
        start = (datetime.now() - timedelta(days=days)).strftime('%Y-%m-%d')
        rows = Reading.query.filter_by(user_id=current_user.id) \
            .filter(Reading.date >= start).all()
        return [r.value for r in rows]

    today_vals = get_vals(0)
    week_vals = get_vals(7)
    month_vals = get_vals(30)

    avg_today = round(statistics.mean(today_vals), 1) if today_vals else 0
    avg_week = round(statistics.mean(week_vals), 1) if week_vals else 0
    avg_month = round(statistics.mean(month_vals), 1) if month_vals else 0
    std_week = round(statistics.stdev(week_vals), 1) if len(week_vals) > 1 else 0
    std_month = round(statistics.stdev(month_vals), 1) if len(month_vals) > 1 else 0

    elements.append(Paragraph("Summary Statistics (Last 30 Days):", subtitle_style))
    stats_data = [
        ["Metric", "Value"],
        ["Today's Average", f"{avg_today}"],
        ["Weekly Average", f"{avg_week}"],
        ["Monthly Average", f"{avg_month}"],
        ["Weekly Std Dev", f"{std_week}"],
        ["Monthly Std Dev", f"{std_month}"],
    ]
    stats_table = Table(stats_data, colWidths=[2.5 * inch, 2 * inch])
    stats_table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#1a73e8')),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
        ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.grey),
        ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#f4f7fb')]),
        ('PADDING', (0, 0), (-1, -1), 8),
    ]))
    elements.append(stats_table)
    elements.append(Spacer(1, 20))

    elements.append(Paragraph("Recent Readings:", subtitle_style))
    readings = Reading.query.filter_by(user_id=current_user.id) \
        .order_by(Reading.date.desc(), Reading.time.desc()) \
        .limit(50).all()

    if readings:
        table_data = [["Date", "Time", "Value", "Meas. Time", "Medication", "Dose", "Carbs(g)"]]  # ✅ تمت الإضافة
        for r in readings:
            table_data.append([r.date, r.time, str(r.value),
                               r.measurement_time or '-', r.medication_name or '-',
                               str(r.medication_dose or 0), str(r.carbs or 0)])  # ✅ تمت الإضافة

        # تعديل عرض الأعمدة عشان تاخد الـ 7 أعمدة
        readings_table = Table(table_data,
                               colWidths=[0.8 * inch, 0.6 * inch, 0.6 * inch, 1.4 * inch, 1.0 * inch, 0.6 * inch,
                                          0.6 * inch])
        readings_table.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#1a73e8')),
            ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
            ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
            ('FONTSIZE', (0, 0), (-1, -1), 9),
            ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
            ('GRID', (0, 0), (-1, -1), 0.5, colors.grey),
            ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#f4f7fb')]),
            ('PADDING', (0, 0), (-1, -1), 5),
        ]))
        elements.append(readings_table)
    else:
        elements.append(Paragraph("No readings recorded.", styles['Normal']))

    doc.build(elements)
    output.seek(0)

    filename = f"diabetes_report_{current_user.username}_{datetime.now().strftime('%Y%m%d')}.pdf"
    return send_file(output, mimetype='application/pdf',
                     as_attachment=True, download_name=filename)